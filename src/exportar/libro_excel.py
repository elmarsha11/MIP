"""
Exporta toda la base a un Excel: los 86 municipios y su evidencia.

Una hoja por tipo de dato, no una hoja por municipio. Con 86 pestañas nadie
compara nada; con una hoja de 86 filas se ordena, se filtra y se manda por mail.

**Cada fila lleva su evidencia.** Es la regla que atraviesa MIP: un numero que no
se puede abrir hasta la fuente no sirve frente a un intendente ni frente a un
auditor. Por eso las hojas de detalle traen la cita textual y la URL, y la hoja
Fuentes explica de donde sale cada cosa y que NO dice.

Uso:
    python src/exportar/libro_excel.py
    python src/exportar/libro_excel.py --salida C:/ruta/MIP.xlsx
"""

from __future__ import annotations

import argparse
import sqlite3
import sys
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Sequence

_AQUI = Path(__file__).resolve().parent
PROJECT_ROOT = _AQUI.parents[1]
for _ruta in (PROJECT_ROOT / "src" / "tablero", PROJECT_ROOT / "src" / "discovery"):
    if str(_ruta) not in sys.path:
        sys.path.insert(0, str(_ruta))

import consultas  # noqa: E402

PROCESADO = PROJECT_ROOT / "data" / "processed"
SALIDA_DIR = PROCESADO / "exportes"

BASES = {
    "oportunidades": PROCESADO / "oportunidades" / "oportunidades_86.sqlite",
    "gabinete": PROCESADO / "gabinete" / "gabinete_86.sqlite",
    "seguridad": PROCESADO / "seguridad" / "seguridad_86.sqlite",
    "operativos": PROCESADO / "seguridad" / "operativos_86.sqlite",
    "territorio": PROCESADO / "territorio" / "entidades_86.sqlite",
    "medios": PROCESADO / "medios" / "medios_86.sqlite",
}

FUENTES = [
    ("Población y sexo", "INDEC, Censo Nacional 2022, resultados definitivos",
     "Donde el Gold Standard difiere, manda INDEC: tiene norma, año y metodología. "
     "Hay 6 municipios cuya población en el Gold Standard es la del partido vecino."),
    ("Autoridades", "Boletín Oficial Municipal (SIBOM) y portal del municipio",
     "Cada nombre sale de una cita literal verificada. La fecha es la del DECRETO, "
     "no la del análisis: es lo que distingue quién ESTÁ de quién ESTUVO."),
    ("Seguridad — cuánto", "SNIC, Ministerio de Seguridad de la Nación, 2025",
     "Son hechos DENUNCIADOS, no delitos ocurridos. El nivel es RELATIVO a los 86: "
     "«alto» significa «en el tercio superior», no «peligroso». El índice excluye "
     "drogas y armas, que se detectan por acción policial y no por denuncia. "
     "Los partidos BALNEARIOS tienen la tasa inflada: los hechos ocurren sobre la "
     "población de verano y se dividen por la residente."),
    ("Seguridad — cómo opera", "Prensa local, últimos 12 meses",
     "Ausencia de evidencia no es evidencia de ausencia: que la prensa leída no lo "
     "haya publicado no prueba que el municipio no lo tenga. 35 municipios no "
     "tienen medios con archivo consultable."),
    ("Oportunidades comerciales", "Portales municipales, releídos buscando fricción",
     "El problema y el nivel de fricción son LECTURA del modelo; lo verificado es "
     "la cita. Requiere ojo humano antes de ir a una propuesta."),
    ("Territorio", "OpenStreetMap, nivel partido (admin_level 5)",
     "OSM es colaborativo y su cobertura varía por municipio. En Chascomús "
     "encuentra 5 de los 7 CAPS que existen: es incompleto y se declara."),
    ("Trámites y gestión digital", "Portal oficial del municipio (Fase 4)",
     "Cada valor con cita literal verificada contra la página. Un vacío se "
     "registra como vacío y no se estima."),
]


def _filas(path: Path, sql: str) -> List[dict]:
    if not path.exists():
        return []
    con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    try:
        return [dict(f) for f in con.execute(sql)]
    except sqlite3.Error:
        return []
    finally:
        con.close()


def hoja_municipios() -> List[dict]:
    """Una fila por municipio: el resumen que se ordena y se filtra."""
    filas = []
    for m in consultas.municipios():
        f = consultas.ficha_resumida(m["municipio"])
        if not f:
            continue
        p, g, c = f["poblacion"], f["seguridad"], f["como_opera"]
        aut = f["autoridades"]
        intendente = aut.get("intendente") or {}
        filas.append({
            "ID": f["id_municipio"],
            "Municipio": f["municipio"],
            "Habitantes": p.get("total"),
            "Mujeres": p.get("mujeres"),
            "% Mujeres": p.get("pct_mujeres"),
            "Varones": p.get("varones"),
            "% Varones": p.get("pct_varones"),
            "Intendente": intendente.get("nombre"),
            "Secretarios": len(aut.get("secretarios") or []),
            "Seguridad - nivel": g.get("nivel") if g.get("disponible") else None,
            "Seguridad - tasa": g.get("tasa_indice") if g.get("disponible") else None,
            "Homicidios": g.get("homicidios") if g.get("disponible") else None,
            "Robos": g.get("robos") if g.get("disponible") else None,
            "Balneario (tasa inflada)": "sí" if g.get("poblacion_estacional") else "",
            "Aspectos de seguridad verificados": sum(
                1 for a in c.get("aspectos", []) if a.get("evidencia")
            ),
            "Hospitales": len(f["salud"].get("hospitales") or []),
            "CAPS": f["salud"].get("caps"),
            "Establecimientos educativos": f["educacion"].get("total"),
            "Canal de turnos": f["salud"].get("turnos_canal"),
            "Sitio oficial": f.get("sitio_oficial"),
            "Sin HTTPS": "sí" if f.get("sitio_sin_https") else "",
        })
    return filas


def hoja_oportunidades() -> List[dict]:
    return [
        {"Municipio": f["municipio"], "Área": f["area"], "Producto UDS": f["producto"],
         "Fricción": f["friccion"], "Situación hoy": f["problema"],
         "Cita textual": f["cita"], "Fuente": f["url"]}
        for f in _filas(
            BASES["oportunidades"],
            "SELECT municipio, area, producto, friccion, problema, cita, url "
            "FROM oportunidades ORDER BY municipio, friccion",
        )
    ]


def hoja_gabinete() -> List[dict]:
    return [
        {"Municipio": f["municipio"], "Cargo": f["cargo"], "Área": f["area"],
         "Nombre": f["nombre"], "Confianza": f["confianza"],
         "Fecha del decreto": f["fecha_norma"], "Fuente": f["fuente"],
         "Cita textual": f["cita"], "URL": f["url"]}
        for f in _filas(
            BASES["gabinete"],
            "SELECT municipio, cargo, area, nombre, confianza, fecha_norma, "
            "fuente, cita, url FROM autoridades ORDER BY municipio, cargo, area",
        )
    ]


def hoja_seguridad() -> List[dict]:
    return [
        {"Municipio": f["municipio"], "Año": f["anio"], "Nivel": f["nivel"],
         "Tasa por 100.000": f["tasa_indice"], "Hechos": f["hechos_indice"],
         "Contra las personas": f["tasa_personas"],
         "Contra la propiedad": f["tasa_propiedad"],
         "Integridad sexual": f["tasa_sexual"],
         "Homicidios dolosos": f["homicidios"], "Robos": f["robos"],
         "Drogas y armas (fuera del índice)": f["tasa_actividad_policial"],
         "Balneario (tasa inflada)": "sí" if f["poblacion_estacional"] else ""}
        for f in _filas(
            BASES["seguridad"],
            "SELECT * FROM seguridad ORDER BY tasa_indice DESC",
        )
    ]


def hoja_como_opera() -> List[dict]:
    return [
        {"Municipio": f["municipio"], "Aspecto": f["aspecto"], "Detalle": f["detalle"],
         "Cita textual": f["cita"], "Medio": f["medio"],
         "Fecha de la nota": f["fecha_nota"], "URL": f["url"]}
        for f in _filas(
            BASES["operativos"],
            "SELECT municipio, aspecto, detalle, cita, medio, fecha_nota, url "
            "FROM operativos ORDER BY municipio, aspecto",
        )
    ]


def hoja_territorio() -> List[dict]:
    return [
        {"Municipio": f["municipio"], "Tipo": f["tipo"], "Nombre": f["nombre"],
         "Dirección": f["direccion"], "Teléfono": f["telefono"], "Fuente": f["url_fuente"]}
        for f in _filas(
            BASES["territorio"],
            "SELECT municipio, tipo, nombre, direccion, telefono, url_fuente "
            "FROM entidades ORDER BY municipio, tipo, nombre",
        )
    ]


def hoja_medios() -> List[dict]:
    return [
        {"Municipio": f["municipio"], "Medio": f["nombre"], "Tipo": f["tipo"],
         "URL": f["url"], "Estado": f["motivo"]}
        for f in _filas(
            BASES["medios"],
            "SELECT municipio, nombre, tipo, url, motivo FROM medios "
            "ORDER BY municipio, tipo, nombre",
        )
    ]


def _escribir(libro, titulo: str, filas: Sequence[dict], anchos: dict = None) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    hoja = libro.create_sheet(titulo)
    if not filas:
        hoja["A1"] = "Sin datos. Corré el módulo correspondiente."
        return

    encabezados = list(filas[0].keys())
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="333333")
        celda.alignment = Alignment(vertical="center")

    for fila in filas:
        hoja.append([fila.get(c) for c in encabezados])

    # Filtro y panel congelado: con 8.478 filas de territorio, sin esto la hoja
    # es inusable.
    hoja.auto_filter.ref = hoja.dimensions
    hoja.freeze_panes = "A2"

    for i, nombre in enumerate(encabezados, start=1):
        largo = max([len(str(nombre))] + [len(str(f.get(nombre) or "")) for f in filas[:200]])
        ancho = (anchos or {}).get(nombre, min(max(largo + 2, 10), 60))
        hoja.column_dimensions[get_column_letter(i)].width = ancho


def _hoja_fuentes(libro) -> None:
    """De donde sale cada dato y —sobre todo— que NO dice.

    Va primera a proposito. Un Excel se reenvia y se lee sin contexto: las
    advertencias tienen que viajar con los numeros o alguien va a leer "alto" en
    seguridad como "peligroso".
    """
    from openpyxl.styles import Alignment, Font

    hoja = libro.create_sheet("Fuentes y advertencias", 0)
    hoja["A1"] = "MIP — Municipal Intelligence Platform"
    hoja["A1"].font = Font(bold=True, size=14)
    hoja["A2"] = f"Exportado el {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    hoja["A3"] = "Confidencial — UDS. No compartir fuera del equipo."
    hoja["A3"].font = Font(bold=True, color="AA0000")

    hoja.append([])
    hoja.append(["Dato", "Fuente", "Qué NO dice"])
    for celda in hoja[5]:
        celda.font = Font(bold=True)
    for dato, fuente, limite in FUENTES:
        hoja.append([dato, fuente, limite])

    for col, ancho in (("A", 26), ("B", 46), ("C", 96)):
        hoja.column_dimensions[col].width = ancho
    for fila in hoja.iter_rows(min_row=6):
        fila[2].alignment = Alignment(wrap_text=True, vertical="top")
        fila[1].alignment = Alignment(wrap_text=True, vertical="top")


def exportar(salida: Optional[Path] = None) -> Path:
    from openpyxl import Workbook

    libro = Workbook()
    libro.remove(libro.active)

    _escribir(libro, "Municipios", hoja_municipios())
    _escribir(libro, "Oportunidades UDS", hoja_oportunidades(),
              {"Situación hoy": 60, "Cita textual": 70, "Fuente": 50})
    _escribir(libro, "Gabinete", hoja_gabinete(), {"Cita textual": 70, "URL": 50})
    _escribir(libro, "Seguridad", hoja_seguridad())
    _escribir(libro, "Seguridad - cómo opera", hoja_como_opera(),
              {"Detalle": 60, "Cita textual": 70, "URL": 50})
    _escribir(libro, "Territorio", hoja_territorio())
    _escribir(libro, "Medios locales", hoja_medios(), {"URL": 45})
    _hoja_fuentes(libro)

    if salida is None:
        SALIDA_DIR.mkdir(parents=True, exist_ok=True)
        salida = SALIDA_DIR / f"MIP_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(salida)
    return salida


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="MIP - Exportar todo a Excel")
    parser.add_argument("--salida", type=Path)
    args = parser.parse_args(argv)

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    ruta = exportar(args.salida)
    tam = ruta.stat().st_size / 1024 / 1024
    print(f"Excel -> {ruta}  ({tam:.1f} MB)")
    print()
    print("Hojas: Fuentes y advertencias · Municipios · Oportunidades UDS ·")
    print("       Gabinete · Seguridad · Seguridad-cómo opera · Territorio · Medios")
    print()
    print("La hoja de advertencias va primera a proposito: un Excel se reenvia y")
    print("se lee sin contexto, asi que los limites viajan con los numeros.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
