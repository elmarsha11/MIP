"""
INDEC - Censo Nacional de Poblacion, Hogares y Viviendas 2022.

Fuente oficial y citable para la poblacion de los 86 municipios. Reemplaza a la
poblacion del Gold Standard, que venia sin fuente ni ano declarados.

    python src/indec/censo.py            muestra el cruce contra el Gold Standard
    python src/indec/censo.py --guardar  escribe data/processed/indec/

Que trae, por partido y con el codigo oficial INDEC:
    poblacion 2022, poblacion 2010, variacion intercensal, superficie, densidad

La apertura por sexo SI esta por partido, en el cuadro 3.2: se creia que solo
existia en REDATAM, pero INDEC publica una hoja por partido dentro del mismo
xlsx. Se incorporo el 2026-08-09.

Lo que todavia NO trae, y se declara faltante en vez de estimarse: la cantidad
de VIVIENDAS por partido. No esta en la serie de poblacion (los cuadros 7 a 9
son indices de envejecimiento); hay que buscarla en la serie de viviendas u
hogares. Es una tarea abierta, no un dato para inventar.
"""

from __future__ import annotations

import io
import json
import re
import sqlite3
import sys
import unicodedata
from pathlib import Path
from typing import Dict, List, NamedTuple, Optional

import requests

_AQUI = Path(__file__).resolve().parent
PROJECT_ROOT = _AQUI.parents[1]
if str(PROJECT_ROOT / "src" / "discovery") not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT / "src" / "discovery"))

SALIDA_DIR = PROJECT_ROOT / "data" / "processed" / "indec"
CACHE_DIR = SALIDA_DIR / "cache"
SQLITE = SALIDA_DIR / "censo_2022.sqlite"

BASE = "https://www.indec.gob.ar/ftp/cuadros/poblacion/"
# Cuadro 2.2: poblacion, superficie y densidad por partido.
# Cuadro 1.2: poblacion 2010 y 2022 por partido, con variacion intercensal.
CUADROS = {
    "densidad": "c2022_bsas_est_c2_2.xlsx",
    "intercensal": "c2022_bsas_est_c1_2.xlsx",
    # Cuadro 3.2: poblacion por sexo registrado al nacer. A diferencia de los
    # otros dos, no es una planilla con una fila por partido: son 136 hojas, una
    # por partido, y el nombre del partido esta en el titulo de cada hoja.
    "sexo": "c2022_bsas_est_c3_2.xlsx",
}
CITA = (
    "INDEC, Censo Nacional de Poblacion, Hogares y Viviendas 2022. "
    "Resultados definitivos, provincia de Buenos Aires."
)

# INDEC usa el nombre completo del partido; el Gold Standard, el de uso corriente.
# Mapeo explicito y verificado uno por uno, nunca fuzzy match: asignarle a un
# municipio la poblacion de otro es exactamente el error que este modulo detecta.
ALIAS_INDEC = {
    "generalmadariaga": "generaljuanmadariaga",
    "sanmigueldelmonte": "monte",
    "veinticincodemayo": "25demayo",
    "alem": "leandronalem",
    "coronelrosales": "coroneldemarinaleonardorosales",
    "gonzaleschaves": "adolfogonzaleschaves",
}


class DatoCenso(NamedTuple):
    codigo_indec: str
    partido: str
    poblacion_2022: Optional[int]
    poblacion_2010: Optional[int]
    variacion_absoluta: Optional[int]
    variacion_relativa: Optional[float]
    superficie_km2: Optional[float]
    densidad: Optional[float]
    mujeres: Optional[int] = None
    varones: Optional[int] = None


def _normalizar(texto: str) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "").replace("\xa0", " ").lower())
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", t)


def _bajar(nombre: str) -> bytes:
    """Descarga con cache en disco: el censo 2022 no va a cambiar."""
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    destino = CACHE_DIR / nombre
    if destino.exists():
        return destino.read_bytes()
    resp = requests.get(
        BASE + nombre, timeout=120, headers={"User-Agent": "MIP-relevamiento-municipal/0.1"}
    )
    resp.raise_for_status()
    if "spreadsheet" not in (resp.headers.get("content-type") or ""):
        raise RuntimeError(f"{nombre}: INDEC no devolvio una planilla")
    destino.write_bytes(resp.content)
    return resp.content


def _hoja(contenido: bytes):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
    hojas = [h for h in wb.sheetnames if h.lower().replace(" ", "").startswith("cuadro")]
    return wb[hojas[0]]


def _numero(v) -> Optional[float]:
    if v is None:
        return None
    try:
        return float(str(v).replace(".", "").replace(",", ".")) if isinstance(v, str) else float(v)
    except ValueError:
        return None


# El titulo de cada hoja es:
#   "Cuadro 3.2.66. Provincia de Buenos Aires, partido Lanus. Total de poblacion, ..."
# El corte va contra ". Total de poblacion", que es fijo, y NO contra el primer
# punto: los partidos con inicial en el nombre ("Jose C. Paz", "Leandro N. Alem")
# quedaban truncados en "Jose C" y se perdian. Alem es uno de los 86.
_RE_TITULO_PARTIDO = re.compile(
    r"partido\s+(.+?)\.\s*Total\s+de\s+poblaci", re.IGNORECASE
)


def leer_sexo() -> Dict[str, tuple]:
    """Poblacion por sexo registrado al nacer, por partido.

    El cuadro 3.2 esta armado distinto de los otros: en vez de una planilla con
    una fila por partido, trae 136 hojas ('Cuadro3.2.0' es la provincia,
    'Cuadro3.2.1' a '.135' los partidos). El nombre del partido no esta en una
    celda de datos sino en el TITULO de cada hoja:

        "Cuadro 3.2.135. Provincia de Buenos Aires, partido Zarate. ..."

    Se lee de ahi y no del indice, que es una hoja aparte: si el orden del indice
    y el de las hojas se desalinearan, cada partido quedaria con la poblacion del
    vecino. Es exactamente el error que este modulo existe para detectar en el
    Gold Standard, asi que no se lo va a reintroducir por comodidad.

    Devuelve {slug: (mujeres, varones, total)}.
    """
    import openpyxl

    wb = openpyxl.load_workbook(
        io.BytesIO(_bajar(CUADROS["sexo"])), data_only=True, read_only=True
    )
    salida: Dict[str, tuple] = {}

    for nombre_hoja in wb.sheetnames:
        if not nombre_hoja.lower().replace(" ", "").startswith("cuadro3.2."):
            continue
        ws = wb[nombre_hoja]
        filas = list(ws.iter_rows(max_row=8, max_col=2, values_only=True))
        if len(filas) < 6:
            continue

        # INDEC separa con espacio duro (\xa0); sin normalizarlo el regex no engancha.
        titulo = next(
            (str(f[0]).replace("\xa0", " ") for f in filas
             if f[0] and "partido" in str(f[0]).lower()),
            "",
        )
        m = _RE_TITULO_PARTIDO.search(titulo)
        if not m:
            continue  # la hoja del total provincial no nombra partido: se saltea
        partido = m.group(1).strip()
        if not partido:
            continue

        total = mujeres = varones = None
        for etiqueta, valor in ((str(f[0] or "").strip().lower(), f[1]) for f in filas):
            if etiqueta == "total":
                total = _numero(valor)
            elif etiqueta.startswith("mujer"):
                mujeres = _numero(valor)
            elif etiqueta.startswith("var"):
                varones = _numero(valor)

        if mujeres is None or varones is None:
            continue

        # Si las partes no suman el total, el cuadro no se leyo como se cree. Un
        # dato mal leido es peor que un dato ausente (ADR-0009).
        if total is not None and abs((mujeres + varones) - total) > 1:
            continue

        salida[_normalizar(partido)] = (int(mujeres), int(varones),
                                        int(total) if total else None)
    return salida


def leer() -> Dict[str, DatoCenso]:
    """Datos del censo por partido, indexados por slug del nombre.

    Se descartan las filas agregadas ('Total', '24 Partidos del Gran Buenos
    Aires', 'Interior'): no son partidos y sumarlas duplicaria poblacion.
    """
    densidad, intercensal = {}, {}

    ws = _hoja(_bajar(CUADROS["densidad"]))
    for fila in ws.iter_rows(min_row=6, values_only=True):
        codigo, partido, sup, pob, dens = (list(fila) + [None] * 5)[:5]
        if not partido or not str(codigo or "").strip().isdigit():
            continue
        densidad[_normalizar(partido)] = (str(codigo).strip(), str(partido).strip(),
                                          _numero(pob), _numero(sup), _numero(dens))

    ws = _hoja(_bajar(CUADROS["intercensal"]))
    for fila in ws.iter_rows(min_row=5, values_only=True):
        codigo, partido, p2010, p2022, var_abs, var_rel = (list(fila) + [None] * 6)[:6]
        if not partido or not str(codigo or "").strip().isdigit():
            continue
        intercensal[_normalizar(partido)] = (_numero(p2010), _numero(var_abs), _numero(var_rel))

    sexo = leer_sexo()

    salida: Dict[str, DatoCenso] = {}
    for clave, (codigo, partido, pob, sup, dens) in densidad.items():
        # Las filas agregadas tienen codigo de 2 digitos ('06'); los partidos, 4 o 5.
        if len(codigo) <= 2:
            continue
        p2010, var_abs, var_rel = intercensal.get(clave, (None, None, None))
        mujeres, varones, total_sexo = sexo.get(clave, (None, None, None))

        # Los dos cuadros tienen que hablar del mismo partido. Si el total por
        # sexo no coincide con la poblacion del cuadro de densidad, se descarta
        # el dato de sexo en vez de mezclar dos partidos distintos.
        if total_sexo is not None and pob is not None and abs(total_sexo - int(pob)) > 1:
            mujeres = varones = None

        salida[clave] = DatoCenso(
            codigo_indec=codigo.zfill(5),
            partido=partido,
            poblacion_2022=int(pob) if pob else None,
            poblacion_2010=int(p2010) if p2010 else None,
            variacion_absoluta=int(var_abs) if var_abs else None,
            variacion_relativa=round(var_rel, 1) if var_rel else None,
            superficie_km2=sup,
            densidad=round(dens, 1) if dens else None,
            mujeres=mujeres,
            varones=varones,
        )
    return salida


def cruzar_con_gold() -> List[dict]:
    """Cruza el censo con los 86 del Gold Standard y muestra las diferencias.

    La comparacion es el punto: el Gold Standard trae poblaciones sin fuente ni
    ano. Donde difieren, manda INDEC, que tiene norma y fecha.
    """
    from discovery_engine import cargar_municipios

    censo = leer()
    filas = []
    for m in cargar_municipios():
        clave = _normalizar(m.nombre)
        d = censo.get(clave) or censo.get(ALIAS_INDEC.get(clave, ""))
        gold = m.poblacion
        indec = d.poblacion_2022 if d else None
        filas.append(
            {
                "municipio": m.nombre,
                "id_municipio": m.id_municipio,
                "codigo_indec": d.codigo_indec if d else None,
                "poblacion_gold": gold,
                "poblacion_indec_2022": indec,
                "poblacion_indec_2010": d.poblacion_2010 if d else None,
                "mujeres": d.mujeres if d else None,
                "varones": d.varones if d else None,
                "variacion_relativa": d.variacion_relativa if d else None,
                "superficie_km2": d.superficie_km2 if d else None,
                "densidad": d.densidad if d else None,
                "diferencia": (indec - gold) if (indec and gold) else None,
                "diferencia_pct": round((indec - gold) / gold * 100, 1)
                if (indec and gold)
                else None,
                "fuente": CITA,
            }
        )
    return filas


def guardar(filas: List[dict]) -> Path:
    SALIDA_DIR.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(SQLITE)
    try:
        con.execute(
            """CREATE TABLE IF NOT EXISTS censo_2022 (
                id_municipio TEXT PRIMARY KEY, municipio TEXT NOT NULL,
                codigo_indec TEXT, poblacion_gold INTEGER, poblacion_indec_2022 INTEGER,
                poblacion_indec_2010 INTEGER, mujeres INTEGER, varones INTEGER,
                variacion_relativa REAL,
                superficie_km2 REAL, densidad REAL, diferencia INTEGER,
                diferencia_pct REAL, fuente TEXT NOT NULL)"""
        )
        # La tabla puede venir de una corrida anterior sin columnas de sexo.
        existentes = {c[1] for c in con.execute("PRAGMA table_info(censo_2022)")}
        for columna in ("mujeres", "varones"):
            if columna not in existentes:
                con.execute(f"ALTER TABLE censo_2022 ADD COLUMN {columna} INTEGER")
        con.executemany(
            "INSERT OR REPLACE INTO censo_2022 (id_municipio,municipio,codigo_indec,"
            "poblacion_gold,poblacion_indec_2022,poblacion_indec_2010,mujeres,varones,"
            "variacion_relativa,superficie_km2,densidad,diferencia,diferencia_pct,fuente) "
            "VALUES (:id_municipio,:municipio,:codigo_indec,:poblacion_gold,"
            ":poblacion_indec_2022,:poblacion_indec_2010,:mujeres,:varones,"
            ":variacion_relativa,:superficie_km2,:densidad,:diferencia,:diferencia_pct,:fuente)",
            filas,
        )
        con.commit()
    finally:
        con.close()
    (SALIDA_DIR / "censo_2022.json").write_text(
        json.dumps({"fuente": CITA, "municipios": filas}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return SQLITE


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    filas = cruzar_con_gold()
    sin_dato = [f for f in filas if not f["poblacion_indec_2022"]]
    con_dif = sorted(
        [f for f in filas if f["diferencia_pct"] is not None and abs(f["diferencia_pct"]) >= 1],
        key=lambda f: -abs(f["diferencia_pct"]),
    )

    print(f"\n{'=' * 76}\nINDEC CENSO 2022 vs GOLD STANDARD\n{'=' * 76}")
    print(f"{CITA}\n")
    print(f"Municipios cruzados : {len(filas) - len(sin_dato)}/{len(filas)}")
    print(f"Poblacion INDEC     : {sum(f['poblacion_indec_2022'] or 0 for f in filas):,}".replace(",", "."))
    print(f"Poblacion Gold      : {sum(f['poblacion_gold'] or 0 for f in filas):,}".replace(",", "."))
    if sin_dato:
        print(f"Sin dato en INDEC   : {[f['municipio'] for f in sin_dato]}")

    print(f"\nDiferencias mayores al 1% ({len(con_dif)} municipios):")
    print(f"  {'Municipio':<24}{'Gold':>10}{'INDEC 2022':>12}{'Dif':>10}{'Dif %':>8}")
    for f in con_dif[:15]:
        print(f"  {f['municipio']:<24}{f['poblacion_gold']:>10,}{f['poblacion_indec_2022']:>12,}"
              f"{f['diferencia']:>10,}{f['diferencia_pct']:>7.1f}%".replace(",", "."))
    print("\nDonde difieren manda INDEC: tiene norma, ano y metodologia publicada.")
    con_sexo = sum(1 for f in filas if f["mujeres"])
    print(f"Apertura por sexo: {con_sexo} de {len(filas)} municipios (cuadro 3.2).")
    print("FALTA (no se estima): viviendas por partido. No esta en la serie de poblacion.")
    print("=" * 76)

    if "--guardar" in sys.argv:
        print(f"\nSQLite -> {guardar(filas)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
